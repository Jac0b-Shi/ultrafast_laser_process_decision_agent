"""Run with the Codex bundled Python after Pandoc conversion."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]/"docs/research/english"
TITLE="AI-Bridged Data-Knowledge Fusion for Intelligent Laser Parameter Recommendation in Ultrafast Laser Processing"
path=ROOT/(TITLE+".docx")
doc=Document(path)
section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=Cm(2);section.bottom_margin=Cm(2);section.left_margin=Cm(2.1);section.right_margin=Cm(2.1)
for name in ('Normal','Body Text','First Paragraph'):
    if name in doc.styles:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(11)
        style.paragraph_format.line_spacing=1.12;style.paragraph_format.space_after=Pt(6)
for name,size in [('Title',17),('Heading 1',13),('Heading 2',11.5),('Caption',9)]:
    if name in doc.styles:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.keep_with_next=True;style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(6)
for p in doc.paragraphs:
    if p.style.name=='Title':
        p.alignment=1
    if re.match(r'^(Figure|Table) \d+\.', p.text):
        p.style=doc.styles['Caption']
        p.paragraph_format.keep_with_next=p.text.startswith('Table ')
    if p._p.xpath('.//w:drawing'):
        p.paragraph_format.keep_with_next=True
    if p.style.name in ('Title', 'Heading 1', 'Heading 2'):
        for run in p.runs:
            run.font.name='Times New Roman'
            run.font.bold=p.style.name != 'Title'
for shape in doc.inline_shapes:
    max_width=Cm(16.4)
    if shape.width>max_width:
        ratio=max_width/shape.width;shape.width=int(shape.width*ratio);shape.height=int(shape.height*ratio)
    max_height=Cm(13.2)
    if shape.height>max_height:
        ratio=max_height/shape.height;shape.width=int(shape.width*ratio);shape.height=int(shape.height*ratio)
for table in doc.tables:
    table.autofit=False
    for row_index,row in enumerate(table.rows):
        properties=row._tr.get_or_add_trPr();properties.append(OxmlElement('w:cantSplit'))
        if row_index==0:properties.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                p.paragraph_format.keep_with_next=row_index<len(table.rows)-1
                for run in p.runs:
                    run.font.name='Times New Roman';run.font.size=Pt(9)
                    if row_index==0:run.bold=True
footer=section.footer.paragraphs[0];footer.alignment=1
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.save(path)
print(path)
