"""Knowledge uploads retain document readers independently of offline authoring tools."""
import io
import pytest
import pymupdf
from docx import Document
from app.services import agent_knowledge as knowledge


@pytest.mark.parametrize('fallback', [False, True])
def test_pdf_upload_extracts_page_evidence(monkeypatch, fallback):
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), 'Pulse density depends on frequency and scanning speed.')
        content = pdf.tobytes()
    if fallback:
        def unavailable_reader(*args, **kwargs):
            raise ValueError('Primary PDF reader failed')
        monkeypatch.setattr(knowledge, 'PdfReader', unavailable_reader)
    chunks = knowledge.extract('knowledge.pdf', content)
    assert chunks[0]['location'] == 'page 1'
    assert 'Pulse density depends on frequency' in chunks[0]['text']


def test_docx_upload_extracts_paragraph_and_table_evidence():
    document = Document()
    document.add_paragraph('Material-specific calibration is required.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Frequency'
    table.cell(0, 1).text = 'kHz'
    content = io.BytesIO()
    document.save(content)
    chunks = knowledge.extract('knowledge.docx', content.getvalue())
    assert chunks == [
        {'location': 'paragraph 1', 'text': 'Material-specific calibration is required.'},
        {'location': 'table 1', 'text': 'Frequency | kHz'},
    ]
